#!/usr/bin/env python3
"""Generate uploadable POS sales-report demos from the shared dataset.

All POS samples under fixtures/pos-sales-reports/ must encode the same
canonical line items from shared/demo-dataset.json.
"""

from __future__ import annotations

import argparse
import json
import sys
from datetime import date, datetime
from pathlib import Path

from openpyxl import Workbook

FIXTURES_ROOT = Path(__file__).resolve().parents[1]
SHARED_DATASET = FIXTURES_ROOT / "shared" / "demo-dataset.json"
ESB_OUTPUT = FIXTURES_ROOT / "esb" / "SalesRecapitulationDetailReport_DEMO.xlsx"
MOKA_OUTPUT = FIXTURES_ROOT / "moka" / "ItemDetails_DEMO.xlsx"

ESB_HEADER_ROW = 11
ESB_DATA_START_ROW = 12
MOKA_HEADER_ROW = 4
MOKA_DATA_START_ROW = 5

# Column order matches ESB Sales Recapitulation Detail Report (SNABB / 2026+).
ESB_HEADERS: list[str] = [
    "Sales Number",
    "Bill Number",
    "Sales Type",
    "Batch Order",
    "Table Section",
    "Table Name",
    "Sales Date",
    "Sales Date In",
    "Sales Date Out",
    "Branch",
    "Brand",
    "City",
    "Area",
    "Visit Purpose",
    "Regular Member Code",
    "Regular Member Name",
    "Loyalty Member Code",
    "Loyalty Member Name",
    "Loyalty Member Type",
    "Employee Code",
    "Employee Name",
    "External Employee Code",
    "External Employee Name",
    "Customer Name",
    "Payment Method",
    "Menu Category",
    "Menu Category Detail",
    "Menu",
    "Custom Menu Name",
    "Menu Code",
    "Menu Notes",
    "Order Mode",
    "Qty",
    "Price",
    "Subtotal",
    "Discount",
    "Service Charge",
    "Tax",
    "VAT",
    "Total",
    "Nett Sales",
    "DPP",
    "Bill Discount",
    "Total After Bill Discount",
    "Waiter",
    "Order Time",
]

# Moka Backoffice → Reports → Transactions → Export Item Details.
# Headers follow Moka's payment + checkout line-item field names (API CheckoutApiItem /
# Payment). Exact vendor Excel labels may vary slightly; this fixture is the
# committed contract for iterative normalizer work until a live dump is available.
MOKA_HEADERS: list[str] = [
    "Payment No",
    "Date",
    "Created At",
    "Outlet Name",
    "Payment Type",
    "Sales Type",
    "Category",
    "Item Name",
    "Variant Name",
    "SKU",
    "Quantity",
    "Price",
    "Gross Sales",
    "Discount",
    "Net Sales",
    "Collected By",
]


def _parse_order_time(value: str) -> datetime:
    return datetime.fromisoformat(value)


def _load_dataset(path: Path) -> dict:
    return json.loads(path.read_text())


def _write_esb(dataset: dict, output_path: Path) -> int:
    meta = dataset["meta"]
    line_items = dataset["line_items"]
    period_start = date.fromisoformat(meta["period_start"])
    period_end = date.fromisoformat(meta["period_end"])

    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "Report"

    ws["A1"] = "Sales Recapitulation Detail Report"
    ws["A2"] = meta["location_name"]
    ws["A4"] = "Generated"
    ws["B4"] = "demo-dataset"
    ws["A5"] = "Period"
    ws["B5"] = (
        f"{period_start.strftime('%d-%m-%Y')} - {period_end.strftime('%d-%m-%Y')}"
    )
    ws["A6"] = "Branch"
    ws["B6"] = "All"
    ws["A7"] = "Sales Type"
    ws["B7"] = meta["sales_type"]
    ws["A8"] = "Generated Username"
    ws["B8"] = "DEMO"
    ws["A9"] = "Report File Name"
    ws["B9"] = "SalesRecapitulationDetailReport_DEMO"

    for col_idx, header in enumerate(ESB_HEADERS, start=1):
        ws.cell(ESB_HEADER_ROW, col_idx, value=header)

    for offset, item in enumerate(line_items):
        order_time = _parse_order_time(item["order_time"])
        sales_date = datetime(order_time.year, order_time.month, order_time.day)
        qty = int(item["qty"])
        price = float(item["price"])
        subtotal = round(price * qty, 2)
        values: dict[str, object] = {
            "Sales Number": item["sales_number"],
            "Bill Number": item["bill_number"],
            "Sales Type": meta["sales_type"],
            "Batch Order": item.get("batch_order", "1"),
            "Table Section": meta["table_section"],
            "Table Name": meta["table_name"],
            "Sales Date": sales_date,
            "Sales Date In": order_time,
            "Sales Date Out": order_time,
            "Branch": meta["location_name"],
            "Brand": meta["brand"],
            "City": meta["city"],
            "Area": meta.get("area"),
            "Visit Purpose": meta["visit_purpose"],
            "Regular Member Code": "Non Member",
            "Regular Member Name": "Non Member",
            "Loyalty Member Code": None,
            "Loyalty Member Name": None,
            "Loyalty Member Type": None,
            "Employee Code": "-",
            "Employee Name": "-",
            "External Employee Code": "-",
            "External Employee Name": "-",
            "Customer Name": "-",
            "Payment Method": item["payment_method"],
            "Menu Category": item["menu_category"],
            "Menu Category Detail": item["menu_category_detail"],
            "Menu": item["menu"],
            "Custom Menu Name": None,
            "Menu Code": item.get("menu_code"),
            "Menu Notes": None,
            "Order Mode": meta["order_mode"],
            "Qty": qty,
            "Price": price,
            "Subtotal": subtotal,
            "Discount": 0.0,
            "Service Charge": 0.0,
            "Tax": 0.0,
            "VAT": 0.0,
            "Total": subtotal,
            "Nett Sales": subtotal,
            "DPP": 0.0,
            "Bill Discount": 0.0,
            "Total After Bill Discount": subtotal,
            "Waiter": "DEMO KASIR",
            "Order Time": order_time,
        }
        excel_row = ESB_DATA_START_ROW + offset
        for col_idx, header in enumerate(ESB_HEADERS, start=1):
            ws.cell(excel_row, col_idx, value=values.get(header))

    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)
    wb.close()
    return len(line_items)


def _write_moka(dataset: dict, output_path: Path) -> int:
    """Encode the shared dataset as a Moka Export Item Details workbook."""
    meta = dataset["meta"]
    line_items = dataset["line_items"]
    period_start = date.fromisoformat(meta["period_start"])
    period_end = date.fromisoformat(meta["period_end"])

    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "Item Details"

    # A1 title used for future POS detection (Reports → Export Item Details).
    ws["A1"] = "Item Details"
    ws["A2"] = meta["location_name"]
    ws["B2"] = (
        f"{period_start.strftime('%d-%m-%Y')} - {period_end.strftime('%d-%m-%Y')}"
    )

    for col_idx, header in enumerate(MOKA_HEADERS, start=1):
        ws.cell(MOKA_HEADER_ROW, col_idx, value=header)

    for offset, item in enumerate(line_items):
        order_time = _parse_order_time(item["order_time"])
        sales_date = datetime(order_time.year, order_time.month, order_time.day)
        qty = int(item["qty"])
        price = float(item["price"])
        subtotal = round(price * qty, 2)
        # Preserve shared category_detail in Variant Name so both POS demos
        # carry identical analytics fields (Moka has no separate subcategory).
        values: dict[str, object] = {
            "Payment No": item["bill_number"],
            "Date": sales_date,
            "Created At": order_time,
            "Outlet Name": meta["location_name"],
            "Payment Type": item["payment_method"],
            "Sales Type": meta["visit_purpose"],
            "Category": item["menu_category"],
            "Item Name": item["menu"],
            "Variant Name": item["menu_category_detail"],
            "SKU": item.get("menu_code"),
            "Quantity": qty,
            "Price": price,
            "Gross Sales": subtotal,
            "Discount": 0.0,
            "Net Sales": subtotal,
            "Collected By": "DEMO KASIR",
        }
        excel_row = MOKA_DATA_START_ROW + offset
        for col_idx, header in enumerate(MOKA_HEADERS, start=1):
            ws.cell(excel_row, col_idx, value=values.get(header))

    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)
    wb.close()
    return len(line_items)


def _smoke_esb(output_path: Path) -> None:
    from menuyukti.core.analytics.esb import normalize_esb_excel
    from menuyukti.core.analytics.pos_detector import detect_pos_from_excel_bytes

    payload = output_path.read_bytes()
    detected = detect_pos_from_excel_bytes(payload)
    if detected != "esb":
        raise SystemExit(f"ERROR: expected POS detect 'esb', got {detected!r}")
    df = normalize_esb_excel(payload)
    if df.empty:
        raise SystemExit("ERROR: normalize_esb_excel returned no rows")
    print(
        f"Smoke ESB: pos={detected} rows={len(df)} "
        f"bills={df['bill_number'].nunique()} menus={df['menu'].nunique()}"
    )


def _smoke_moka(output_path: Path, expected_rows: int) -> None:
    import pandas as pd
    from openpyxl import load_workbook

    wb = load_workbook(output_path, data_only=True)
    ws = wb.active
    assert ws is not None
    a1 = ws["A1"].value
    wb.close()
    if a1 != "Item Details":
        raise SystemExit(f"ERROR: expected Moka A1 'Item Details', got {a1!r}")

    df = pd.read_excel(output_path, skiprows=MOKA_HEADER_ROW - 1)
    missing = [h for h in MOKA_HEADERS if h not in df.columns]
    if missing:
        raise SystemExit(f"ERROR: Moka fixture missing columns: {missing}")
    if len(df) != expected_rows:
        raise SystemExit(f"ERROR: expected {expected_rows} Moka rows, got {len(df)}")
    if df["Payment No"].nunique() != 4:
        raise SystemExit("ERROR: expected 4 distinct Moka Payment No values")
    print(
        f"Smoke Moka: a1={a1!r} rows={len(df)} "
        f"bills={df['Payment No'].nunique()} menus={df['Item Name'].nunique()}"
    )


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(
        description="Generate POS sales-report demos from shared/demo-dataset.json"
    )
    parser.add_argument(
        "--dataset",
        type=Path,
        default=SHARED_DATASET,
        help=f"Canonical shared dataset (default: {SHARED_DATASET})",
    )
    parser.add_argument(
        "--pos",
        choices=("esb", "moka", "all"),
        default="all",
        help="Which POS demo(s) to generate (default: all)",
    )
    parser.add_argument(
        "--skip-smoke",
        action="store_true",
        help="Skip menuyukti normalizer / fixture smoke check",
    )
    args = parser.parse_args(argv)

    dataset = _load_dataset(args.dataset)
    targets: list[str] = []
    if args.pos == "all":
        targets = ["esb", "moka"]
    else:
        targets = [args.pos]

    if "esb" in targets:
        n = _write_esb(dataset, ESB_OUTPUT)
        print(f"Wrote {ESB_OUTPUT} ({n} line items)")
        if not args.skip_smoke:
            _smoke_esb(ESB_OUTPUT)

    if "moka" in targets:
        n = _write_moka(dataset, MOKA_OUTPUT)
        print(f"Wrote {MOKA_OUTPUT} ({n} line items)")
        if not args.skip_smoke:
            _smoke_moka(MOKA_OUTPUT, expected_rows=n)

    return 0


if __name__ == "__main__":
    sys.exit(main())
