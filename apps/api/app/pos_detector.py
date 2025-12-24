from typing import Optional
from openpyxl import load_workbook
from io import BytesIO


def detect_pos_from_excel_bytes(data: bytes) -> Optional[str]:
    """
    Detect POS by reading cell A1 from the first worksheet.
    """

    wb = load_workbook(filename=BytesIO(data), read_only=True, data_only=True)
    ws = wb.worksheets[0]

    value = ws["A1"].value

    if not isinstance(value, str):
        return None

    normalized = value.strip().lower()

    if "sales recapitulation detail report" in normalized:
        return "esb"

    return "unknown"
